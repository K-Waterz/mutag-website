"""Generate Mutag House Company Profile (Word) for tender use."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parent / "Mutag-House-Company-Profile.docx"

BLUE = RGBColor(0x1E, 0x3A, 0x5F)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4B, 0x55, 0x63)

ADDRESS = "2957B Goshawk Street, Thatchfield Hills, Centurion, 0157"
REG_NUMBER = "2025/500507/07"


def set_run(run, *, size=11, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    size = {1: 16, 2: 13, 3: 11}.get(level, 11)
    set_run(run, size=size, bold=True, color=BLUE)
    return p


def add_body(doc, text, *, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, size=11, bold=bold, color=DARK)
    return p


def add_bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run(r1, size=11, bold=True, color=DARK)
        r2 = p.add_run(text)
        set_run(r2, size=11, color=DARK)
    else:
        run = p.add_run(text)
        set_run(run, size=11, color=DARK)
    return p


def set_cell_text(cell, text, *, bold=False, size=10, color=DARK, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_info_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    for i, (label, value) in enumerate(rows):
        set_cell_text(table.rows[i].cells[0], label, bold=True, size=10, color=BLUE)
        set_cell_text(table.rows[i].cells[1], value, size=10)
        table.rows[i].cells[0].width = Inches(1.8)
        table.rows[i].cells[1].width = Inches(4.7)
    return table


def add_client_entry(doc, client, work_items):
    """Client name as bold heading-style bullet intro, then work list."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(client)
    set_run(run, size=11, bold=True, color=BLUE)

    for item in work_items:
        add_bullet(doc, item)


def build():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # Cover / title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("MUTAG HOUSE (PTY) LTD")
    set_run(r, size=22, bold=True, color=BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    r = subtitle.add_run("Company Profile")
    set_run(r, size=16, bold=True, color=ACCENT)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag.paragraph_format.space_after = Pt(4)
    r = tag.add_run("Software Development Studio & Digital Systems Company")
    set_run(r, size=11, color=MUTED)

    prepared = doc.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_after = Pt(14)
    r = prepared.add_run("August 2026")
    set_run(r, size=10, color=MUTED)

    # 1. Overview
    add_heading(doc, "1. Company Overview", 1)
    add_body(
        doc,
        "Mutag House is a software development studio specialising in the design and "
        "delivery of scalable digital systems. Based in Centurion, Gauteng, we build "
        "the platforms, applications, and backend infrastructure that organisations "
        "rely on to operate, grow, and scale.",
    )
    add_body(
        doc,
        "Our work goes beyond websites. We architect and implement custom web "
        "applications, secure authentication systems, APIs, dashboards, and automated "
        "workflows that support real business processes and complex operational "
        "requirements. Every system we build is grounded in clean architecture, clear "
        "business logic, and long-term maintainability.",
    )
    add_body(
        doc,
        "We work with startups, SMEs, corporates, and institutions that require systems "
        "that integrate with existing tools, enforce access control, manage data "
        "securely, and perform reliably at scale. Branding, digital presence, and "
        "business compliance, and ICT services support the systems we build — our focus remains "
        "on engineering digital infrastructure that works, lasts, and delivers "
        "measurable business impact.",
    )
    add_body(
        doc,
        "At our core, we are builders of systems — practical, scalable, and designed "
        "for the realities of modern business.",
        bold=True,
    )

    add_heading(doc, "Company Particulars", 2)
    add_info_table(
        doc,
        [
            ("Legal name", "Mutag House (Pty) Ltd"),
            ("Trading as", "Mutag House"),
            ("Registration number", REG_NUMBER),
            ("Physical address", ADDRESS),
            ("Website", "www.mutag.co.za"),
            ("Email", "info@mutag.co.za"),
            ("Telephone", "+27 72 957 2238"),
            ("Business hours", "Monday–Friday, 08:00–17:00"),
            (
                "Service area",
                "South Africa (Gauteng-focused; remote nationwide and international)",
            ),
        ],
    )

    # 2–4 Mission / Vision / Values
    add_heading(doc, "2. Mission", 1)
    add_body(
        doc,
        "To design and deliver scalable software systems that help South African "
        "businesses and institutions operate with clarity, efficiency, and confidence "
        "— combining technical excellence with practical business understanding.",
    )

    add_heading(doc, "3. Vision", 1)
    add_body(
        doc,
        "To be a trusted digital systems partner for growing organisations — known for "
        "clean architecture, reliable delivery, and solutions that endure beyond launch.",
    )

    add_heading(doc, "4. Core Values", 1)
    add_bullet(
        doc,
        " We hold every project to high standards of quality and professionalism.",
        bold_prefix="Excellence —",
    )
    add_bullet(
        doc,
        " We stay current with proven technologies and design approaches that deliver lasting value.",
        bold_prefix="Innovation —",
    )
    add_bullet(
        doc,
        " We treat every client relationship as a collaboration, not a transaction.",
        bold_prefix="Partnership —",
    )
    add_bullet(
        doc,
        " We commit to on-time delivery, clear communication, and systems that perform under real-world conditions.",
        bold_prefix="Reliability —",
    )

    # 5 Services
    add_heading(doc, "5. Services Offered", 1)
    add_body(
        doc,
        "Mutag House organises its work into five complementary service areas.",
    )

    add_heading(doc, "5.1 Software Engineering (Core)", 2)
    add_body(doc, "Production-grade systems that run the business:")
    add_bullet(
        doc,
        " Progressive Web Apps (PWAs), Single Page Applications (SPAs), multi-tenant SaaS platforms, and enterprise portals",
        bold_prefix="Custom Web Application Development —",
    )
    add_bullet(
        doc,
        " Payment gateway integrations, inventory management, order processing, and subscription billing",
        bold_prefix="E-Commerce & Transaction Platforms —",
    )
    add_bullet(
        doc,
        " RESTful and GraphQL APIs, database design, cloud infrastructure, and third-party integrations",
        bold_prefix="Backend Systems & APIs —",
    )
    add_bullet(
        doc,
        " Real-time analytics, data visualisation, KPI tracking, and reporting tools",
        bold_prefix="Business Intelligence & Dashboards —",
    )
    add_bullet(
        doc,
        " Workflow engines, CRM/ERP integrations, automated notifications, and custom automation tooling",
        bold_prefix="Automation & Workflow Systems —",
    )
    add_bullet(
        doc,
        " OAuth, SSO, role-based access control (RBAC), API security, and POPIA/GDPR-aligned compliance",
        bold_prefix="Authentication & Security Systems —",
    )

    add_heading(doc, "5.2 Web & Digital Presence", 2)
    add_bullet(doc, "Corporate websites and conversion-focused landing pages")
    add_bullet(doc, "SEO-optimised web properties and ongoing website maintenance")
    add_bullet(
        doc,
        "Growth and go-to-market support (email campaigns, SEO, content strategy) where it complements software delivery",
    )

    add_heading(doc, "5.3 Branding & Design", 2)
    add_bullet(doc, "Logo design, brand guidelines, and corporate identity packages")
    add_bullet(
        doc,
        "Marketing materials and print services (business cards, stationery, large-format)",
    )
    add_bullet(doc, "Corporate giftings and related print production")

    add_heading(doc, "5.4 Business & Compliance", 2)
    add_bullet(doc, "Company registration and business documentation")
    add_bullet(doc, "CIPC beneficial ownership filings")
    add_bullet(
        doc,
        "CRM implementation and presentation design (investor decks, sales materials)",
    )

    add_heading(doc, "5.5 ICT & Infrastructure", 2)
    add_body(
        doc,
        "Reliable ICT that keeps the business secure and connected. Cybersecurity sits at the centre of every engagement we design.",
    )
    add_bullet(
        doc,
        " Endpoint monitoring, firewalls, MFA, backup and disaster recovery, biometric/RFID access, smart locks, CCTV, and security audits",
        bold_prefix="Cybersecurity & Network Protection —",
    )
    add_bullet(
        doc,
        " Remote and on-site support with SLAs, device monitoring, software and licence management, and 24/7 support",
        bold_prefix="Managed IT Services —",
    )
    add_bullet(
        doc,
        " Microsoft 365 or Google Workspace migration, cloud backup, secure storage, and real-time collaboration tools",
        bold_prefix="Cloud Solutions —",
    )
    add_bullet(
        doc,
        " Structured cabling (Cat5e, Cat6, fibre), server rooms and racks, WiFi design, UPS/backup power, and IoT hardware",
        bold_prefix="Infrastructure Services —",
    )
    add_bullet(
        doc,
        " Multi-floor network installations, fleet-wide endpoint deployments, digital transformation, and scoped solutions",
        bold_prefix="Custom ICT Projects —",
    )

    # 6 Approach
    add_heading(doc, "6. Our Approach", 1)
    add_bullet(
        doc,
        " We design software with scalability, maintainability, and clean architecture at the centre — systems built to grow, integrate with existing infrastructure, and remain operable over time.",
        bold_prefix="Systems Architecture —",
    )
    add_bullet(
        doc,
        " We translate complex operational requirements into robust, reliable software so that business rules are accurately reflected in code.",
        bold_prefix="Business Logic Engineering —",
    )
    add_bullet(
        doc,
        " From frontend interfaces to backend APIs, databases, authentication, and automation — we cover the full stack to ensure coherent performance across the system.",
        bold_prefix="Full-Stack Capability —",
    )

    # 7 Process
    add_heading(doc, "7. Development Process", 1)
    add_bullet(
        doc,
        " Understand business requirements and design a technical architecture that can scale.",
        bold_prefix="1. Discovery & System Architecture —",
    )
    add_bullet(
        doc,
        " Translate operational rules into maintainable, testable software.",
        bold_prefix="2. Business Logic Mapping —",
    )
    add_bullet(
        doc,
        " Build in cycles, test continuously, and refine against real feedback.",
        bold_prefix="3. Iterative Development & Testing —",
    )
    add_bullet(
        doc,
        " Launch smoothly, monitor performance, and provide ongoing maintenance and scaling support.",
        bold_prefix="4. Deployment, Scaling & Long-Term Support —",
    )

    # 8 Industries
    add_heading(doc, "8. Industries Served", 1)
    add_body(
        doc,
        "Mutag House works across sectors that need scalable digital systems, including "
        "fintech, healthcare, logistics, retail and e-commerce, education, professional "
        "services (including legal), and hospitality and lifestyle brands. Solutions are "
        "adaptable to any organisation requiring custom platforms, secure data handling, "
        "or process automation.",
    )

    # 9 Selected Client Work
    add_heading(doc, "9. Selected Client Work", 1)
    add_body(
        doc,
        "These are some of the companies we have worked with, and the nature of the "
        "work completed for each:",
    )

    # Website portfolio + track-record additions merged (no contacts / no amounts)
    add_client_entry(
        doc,
        "XFM Tech",
        [
            "Website design and development",
            "Brand identity",
            "Search engine optimisation (SEO)",
            "Custom web application development",
            "Social media marketing and branded visual content",
        ],
    )
    add_client_entry(
        doc,
        "FlowerClub",
        [
            "Social media transformation and brand storytelling",
            "Social media strategy and content creation",
            "Digital marketing strategy",
            "Paid advertising",
            "Ongoing website maintenance",
        ],
    )
    add_client_entry(
        doc,
        "Mathibele Mahlaela Attorneys",
        [
            "Website design and development",
            "Legal branding solutions and professional business cards",
            "Corporate giftings and printing services",
            "CRM implementation",
            "CIPC compliance and regulatory filing / business documentation",
        ],
    )
    add_client_entry(
        doc,
        "Magnificent Pools",
        [
            "Corporate branding and graphic design",
            "Lead-capture landing page campaign for lead generation",
            "Lead capture system with database integration",
            "Social media visual advertising",
        ],
    )
    add_client_entry(
        doc,
        "We Moove SA",
        [
            "Custom internal web-based quoting system",
            "Internal tools to streamline pricing workflows and improve turnaround time",
        ],
    )
    add_client_entry(
        doc,
        "Rapatla Rue Attorneys",
        [
            "Complete digital business setup and advertising",
            "Custom website development",
            "Premium hosting setup and professional email configuration",
            "SEO optimisation",
            "Digital advertising campaigns",
        ],
    )
    add_client_entry(
        doc,
        "La Gracia Estate",
        [
            "Complete brand identity (logo, business cards, and letterhead)",
            "Strategic brand planning and cohesive stationery package",
        ],
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(8)
    r = note.add_run(
        "Further case studies and project details are available at www.mutag.co.za/portfolio."
    )
    set_run(r, size=9, color=MUTED)

    # 10 Why Mutag
    add_heading(doc, "10. Why Mutag House", 1)
    add_bullet(
        doc,
        " Based in Centurion with familiarity with South African business, regulatory, and market realities (including POPIA).",
        bold_prefix="Local expertise —",
    )
    add_bullet(
        doc,
        " We build infrastructure and applications that run operations, not only marketing surfaces.",
        bold_prefix="Systems-first mindset —",
    )
    add_bullet(
        doc,
        " Software, web presence, branding, and compliance can be coordinated under one partner when required.",
        bold_prefix="Integrated delivery —",
    )
    add_bullet(
        doc,
        " Solutions are scoped to each client’s goals, constraints, and growth stage.",
        bold_prefix="Personalised engagement —",
    )
    add_bullet(
        doc,
        " We remain available for maintenance, iteration, and scaling after go-live.",
        bold_prefix="Long-term support —",
    )

    # 11 Contact
    add_heading(doc, "11. Contact Details", 1)
    add_info_table(
        doc,
        [
            ("Company", "Mutag House (Pty) Ltd"),
            ("Registration number", REG_NUMBER),
            ("Physical address", ADDRESS),
            ("Website", "www.mutag.co.za"),
            ("Email", "info@mutag.co.za"),
            ("Telephone", "+27 72 957 2238"),
            ("WhatsApp", "wa.link/rgwahs"),
            ("LinkedIn", "linkedin.com/company/mutag-house"),
            ("Instagram", "@mutag_house"),
            ("Facebook", "MUTAG House"),
        ],
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(24)
    r = footer.add_run(
        "© 2026 Mutag House (Pty) Ltd. All rights reserved.\n"
        "Company Profile — August 2026"
    )
    set_run(r, size=9, color=MUTED)

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
